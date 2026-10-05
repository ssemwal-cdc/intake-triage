"""Self-check for build_slide.py: no framework, just asserts.

Run: python template/selfcheck_build_slide.py  (from skills/request-slide/)
"""
import json
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
import build_slide  # noqa: E402

from pptx import Presentation

SAMPLE = Path(__file__).resolve().parent.parent / 'examples' / 'sample.json'


def shape_geometry(path):
    prs = Presentation(str(path))
    slide = prs.slides[0]
    return [(s.shape_type, s.left, s.top, s.width, s.height) for s in slide.shapes]


def check_idempotent():
    fields = json.loads(SAMPLE.read_text(encoding='utf-8'))
    with tempfile.TemporaryDirectory() as d:
        p1 = Path(d) / 'a.pptx'
        p2 = Path(d) / 'b.pptx'
        build_slide.build(fields, p1)
        build_slide.build(fields, p2)
        g1, g2 = shape_geometry(p1), shape_geometry(p2)
        assert g1 == g2, 'shape positions/sizes differ between two runs on the same JSON'
    print('ok: identical shape geometry across two runs')


def check_overflow_fails_loudly():
    fields = json.loads(SAMPLE.read_text(encoding='utf-8'))
    fields['ask'] = ' '.join(['word'] * 31)
    try:
        build_slide.check_budgets(fields)
    except ValueError as e:
        assert 'ask' in str(e)
        print('ok: over-budget ask raises ValueError')
        return
    raise AssertionError('expected ValueError for over-budget ask')


if __name__ == '__main__':
    check_idempotent()
    check_overflow_fails_loudly()
    print('all checks passed')
